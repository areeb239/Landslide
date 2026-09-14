import os
import shutil
import subprocess
import pymupdf
from pptx import Presentation
from pptx.util import Inches

edge = r'C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe'
pdf_out = r'c:\Users\HP\.gemini\antigravity-ide\scratch\Landslide\sih_presentation\Bhoochetak_SIH_2026.pdf'
html_url = r'file:///c:/Users/HP/.gemini/antigravity-ide/scratch/Landslide/sih_presentation/index.html'

print("1. Rendering PDF with Edge headless...")
res = subprocess.run([
    edge,
    '--headless=new',
    '--disable-gpu',
    '--no-pdf-header-footer',
    f'--print-to-pdf={pdf_out}',
    html_url
], capture_output=True, text=True)

print("Edge status:", res.returncode)

print("2. Rendering slide PNG previews with PyMuPDF...")
doc = pymupdf.open(pdf_out)
print(f"Total pages: {len(doc)}")
for i, page in enumerate(doc):
    pix = page.get_pixmap(dpi=150)
    out_img = f"sih_presentation/slide_{i+1}.png"
    pix.save(out_img)
    print(f"Rendered {out_img} ({pix.width}x{pix.height})")

print("3. Building 16:9 Widescreen PPTX...")
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

for i in range(1, 7):
    img_path = f"sih_presentation/slide_{i}.png"
    if os.path.exists(img_path):
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(img_path, 0, 0, width=Inches(13.333), height=Inches(7.5))

pptx_out = r"sih_presentation\Bhoochetak_SIH_2026.pptx"
prs.save(pptx_out)
print(f"Saved PPTX to: {pptx_out}")

print("4. Copying files to project root...")
shutil.copy2(pdf_out, r"c:\Users\HP\.gemini\antigravity-ide\scratch\Landslide\Bhoochetak_SIH_2026.pdf")
shutil.copy2(pptx_out, r"c:\Users\HP\.gemini\antigravity-ide\scratch\Landslide\Bhoochetak_SIH_2026.pptx")
print("All artifacts generated and copied to root successfully!")
